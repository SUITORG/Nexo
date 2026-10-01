# -*- coding: utf-8 -*-
r"""Motor del SuitCVLO-open.

Pipeline (portado del proceso validado):
  fotos (web/local/zip/carpeta) -> GPS EXIF -> match a PIN (radio) -> lectura IA
  -> copia a <SALIDA>\<AAMMPQ>\auditoria\<Marca><Tipo><Ciudad>\<PIN>\ (HEIC->JPG)
  -> info.txt por PIN (foto -> pines -> KML) -> idxfotos.xlsx (crear/actualizar, Verificado).

Sin base de datos: todo vive en archivos locales. Sin operaciones destructivas:
nunca se elimina ningún archivo (los temporales se quedan en trabajo\\cache).

MODO_MOCK (env CVLO_MOCK=1): sin red; para pruebas.
"""
import hashlib
import json
import math
import mimetypes
import os
import re
import time
import urllib.request
import uuid
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor

import pillow_heif
from PIL import Image
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import utilidades as U

pillow_heif.register_heif_opener()


def en_mock():
    return os.environ.get('CVLO_MOCK') == '1'


# ------------------------------------------- API AutoGLM (último eslabón de la cadena)
# Sin claves en el código: APP_ID/APP_KEY por .env (AUTOCALW_*) o desde la skill instalada.
DIR_SKILL = os.environ.get('CVLO_SKILL_DIR', r'C:\Users\rojo-\.openclaw-autoclaw\skills\autoglm-image-recognition')
CONFIG_DEFECTO = {
    'APP_ID': '',
    'APP_KEY': '',
    'TOKEN_URL': 'http://127.0.0.1:18432/get_token',
    'URL_UPLOAD': 'https://autoglm-api.autoglm.ai/agentdr/v1/assistant/upload-mix',
    'URL_RECONOCER': 'https://autoglm-api.autoglm.ai/agentdr/v1/assistant/skills/image-recognition',
}
ENTORNO_CFG = {'APP_ID': 'AUTOCALW_APP_ID', 'APP_KEY': 'AUTOCALW_APP_KEY',
               'TOKEN_URL': 'AUTOCALW_TOKEN_URL', 'URL_UPLOAD': 'AUTOCALW_URL_UPLOAD',
               'URL_RECONOCER': 'AUTOCALW_URL_RECONOCER'}


def leer_config():
    try:
        from ia import switch
        switch._env('AUTOCALW_APP_KEY')   # efecto secundario: carga el .env del proyecto
    except Exception:
        pass
    cfg = dict(CONFIG_DEFECTO)
    try:
        texto = open(os.path.join(DIR_SKILL, 'upload-mix.py'), encoding='utf-8').read()

        def valor(clave):
            m = re.search(r'^\s*{}\s*=\s*"([^"]+)"'.format(clave), texto, re.M)
            return m.group(1) if m else None

        for clave in ('APP_ID', 'APP_KEY', 'TOKEN_URL'):
            v = valor(clave)
            if v:
                cfg[clave] = v
        v = valor('URL')
        if v:
            cfg['URL_UPLOAD'] = v
    except Exception:
        pass
    try:
        texto = open(os.path.join(DIR_SKILL, 'image-recognition.py'), encoding='utf-8').read()
        m = re.search(r'^\s*URL\s*=\s*"([^"]+)"', texto, re.M)
        if m:
            cfg['URL_RECONOCER'] = m.group(1)
    except Exception:
        pass
    # el entorno (.env) manda sobre cualquier valor leído
    for clave, var in ENTORNO_CFG.items():
        v = os.environ.get(var)
        if v:
            cfg[clave] = v
    if not cfg.get('APP_KEY'):
        raise RuntimeError('Falta AUTOCALW_APP_KEY (AutoClaw es el último eslabón; ver .env.example)')
    return cfg


def token_local(timeout=6):
    cfg = leer_config()
    with urllib.request.urlopen(cfg['TOKEN_URL'], timeout=timeout) as r:
        token = r.read().decode('utf-8').strip()
    if not token.lower().startswith('bearer '):
        token = 'Bearer ' + token
    return token, cfg


def token_disponible():
    if en_mock():
        return True
    try:
        token_local(timeout=3)
        return True
    except Exception:
        return False


def subir_imagen(ruta, reintentos=2):
    ultimo = None
    for intento in range(reintentos + 1):
        try:
            token, cfg = token_local()
            ts = str(int(time.time()))
            firma = hashlib.md5('{}&{}&{}'.format(cfg['APP_ID'], ts, cfg['APP_KEY']).encode('utf-8')).hexdigest()
            nombre = os.path.basename(ruta)
            mime = mimetypes.guess_type(ruta)[0] or 'application/octet-stream'
            datos = open(ruta, 'rb').read()
            boundary = '----CVLOBoundary' + uuid.uuid4().hex[:16]
            parte = ('--{b}\r\nContent-Disposition: form-data; name="files"; filename="{n}"\r\nContent-Type: {m}\r\n\r\n'
                     .format(b=boundary, n=nombre, m=mime)).encode('utf-8')
            cuerpo = parte + datos + ('\r\n--{}--\r\n'.format(boundary)).encode('utf-8')
            headers = {'Authorization': token, 'Content-Type': 'multipart/form-data; boundary=' + boundary,
                       'X-Auth-Appid': cfg['APP_ID'], 'X-Auth-TimeStamp': ts, 'X-Auth-Sign': firma}
            req = urllib.request.Request(cfg['URL_UPLOAD'], data=cuerpo, headers=headers, method='POST')
            with urllib.request.urlopen(req, timeout=180) as resp:
                j = json.loads(resp.read().decode('utf-8'))
            return j['data']['oss_info'][0]['oss_url']
        except Exception as e:
            ultimo = e
            time.sleep(2 + intento * 3)
    raise RuntimeError('No se pudo subir la imagen ({}). Verifica AutoClaw abierto e internet.'.format(ultimo))


def reconocer(url_imagen, prompt, reintentos=2, timeout=300):
    ultimo = None
    for intento in range(reintentos + 1):
        try:
            token, cfg = token_local()
            ts = str(int(time.time()))
            firma = hashlib.md5('{}&{}&{}'.format(cfg['APP_ID'], ts, cfg['APP_KEY']).encode('utf-8')).hexdigest()
            headers = {'Authorization': token, 'Content-Type': 'application/json',
                       'X-Auth-Appid': cfg['APP_ID'], 'X-Auth-TimeStamp': ts, 'X-Auth-Sign': firma}
            payload = json.dumps({'prompt': prompt, 'image_url': url_imagen}).encode('utf-8')
            req = urllib.request.Request(cfg['URL_RECONOCER'], data=payload, headers=headers, method='POST')
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                j = json.loads(resp.read().decode('utf-8'))
            texto = (j.get('data') or {}).get('text')
            if texto:
                return texto
            raise RuntimeError('Respuesta sin texto')
        except Exception as e:
            ultimo = e
            time.sleep(3 + intento * 3)
    raise RuntimeError('No se pudo leer la imagen: {}'.format(ultimo))


PROMPT_LECTURA = ('Responde únicamente en ESPAÑOL, con este formato exacto de 7 líneas:\n'
                  'ANUNCIO: SI o NO (¿se ve algún anuncio publicitario o propaganda en la foto?)\n'
                  'MARCA CAMPAÑA: ... (marca, partido, evento o campaña visible; si no, \'no visible\')\n'
                  'QUE ANUNCIA: ... (1-2 frases de qué promociona)\n'
                  'TEXTO VISIBLE: ... (transcripción literal de lo principal)\n'
                  'REFERENCIAS UBICACION: ... (calles, negocios, códigos visibles o \'no visible\')\n'
                  'TIPO SOPORTE: ... (panorámico, muro, parabús, mupi, espectacular, cartelera, valla u otro)\n'
                  'CARA: frente/reverso/ambos/no se distingue\n'
                  'Sé breve.')


def campo(texto, nombre):
    for l in (texto or '').splitlines():
        l = l.strip()
        if l.upper().startswith(nombre.upper()) and ':' in l:
            return l.split(':', 1)[1].strip()
    return ''


# ---------------------------------------------------------------- GPS / fechas
def gps_y_fecha(ruta, mock_gps=None):
    """Devuelve (lat, lon, fecha 'YYYY-MM-DD' o None). Corrige hemisferios S/W."""
    nombre = os.path.basename(ruta)
    if en_mock() and mock_gps is not None:
        d = mock_gps.get(nombre)
        if not d:
            return (None, None, None)
        return (d.get('lat'), d.get('lon'), d.get('fecha'))
    try:
        im = Image.open(ruta)
        exif = im.getexif()
        fecha = None
        f = exif.get(36867) or exif.get(306)
        if f:
            m = re.match(r'(\d{4})[:\-](\d{2})[:\-](\d{2})', str(f))
            if m:
                fecha = '{}-{}-{}'.format(m.group(1), m.group(2), m.group(3))
        gps = exif.get_ifd(34853)
        if not gps:
            return (None, None, fecha)

        def g(v):
            try:
                return float(v[0]) + float(v[1]) / 60 + float(v[2]) / 3600
            except Exception:
                try:
                    return float(v)
                except Exception:
                    return None

        lat = g(gps.get(2))
        lon = g(gps.get(4))
        if lat is None or lon is None:
            return (None, None, fecha)
        if str(gps.get(1)).upper().startswith('S'):
            lat = -lat
        if str(gps.get(3)).upper().startswith('W'):
            lon = -lon
        return (lat, lon, fecha)
    except Exception:
        return (None, None, None)


def dist_m(lat1, lon1, lat2, lon2):
    r = 6371000.0
    p1 = math.radians(lat1)
    p2 = math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


def quincena_de(fecha_iso):
    """'2026-09-14' -> '260901' (día<=15 -> 01; >=16 -> 02). Sin fecha -> None."""
    if not fecha_iso:
        return None
    try:
        y, m, d = [int(x) for x in fecha_iso.split('-')[:3]]
        pq = '01' if d <= 15 else '02'
        return '{:02d}{:02d}{}'.format(y % 100, m, pq)
    except Exception:
        return None


# ---------------------------------------------------------------- pines
PRIM_HEADERS = ['PIN', 'LATITUD', 'LONGITUD', 'MEDIO', 'FORMATO', 'ID', 'UBICACIÓN', 'CIUDAD',
                'DIRECCIÓN', 'MARCA', 'CLIENTE', 'CODIGO MUPI', 'CAPA', 'FAMILIA']


def _pines_desde_kml(ruta_kml):
    import xml.etree.ElementTree as ET
    NS = '{http://www.opengis.net/kml/2.2}'
    data = open(ruta_kml, 'rb').read()
    raiz = ET.fromstring(data)

    def nombre_de(el):
        n = el.find(NS + 'name')
        return (n.text or '').strip() if n is not None and n.text else ''

    pines = {}

    def walk(el, capa):
        for hijo in el:
            tag = hijo.tag.split('}')[-1]
            if tag == 'Placemark':
                datos = {}
                for d in hijo.findall('.//' + NS + 'Data'):
                    v = d.find(NS + 'value')
                    datos[d.get('name', '')] = (v.text or '').strip() if v is not None else ''
                pin = nombre_de(hijo)
                if not pin:
                    continue
                d = {
                    'PIN': pin, 'CAPA': capa,
                    'LATITUD': datos.get('LATITUD', ''), 'LONGITUD': datos.get('LONGITUD', ''),
                    'MEDIO': datos.get('MEDIO', '') or datos.get('TIPO DE MEDIO', ''),
                    'FORMATO': datos.get('FORMATO', ''), 'ID': datos.get('ID', ''),
                    'UBICACIÓN': datos.get('UBICACIÓN', ''), 'CIUDAD': datos.get('CIUDAD', ''),
                    'DIRECCIÓN': datos.get('DIRECCIÓN', ''), 'MARCA': datos.get('MARCA', ''),
                    'CLIENTE': datos.get('CLIENTE', ''), 'CODIGO MUPI': datos.get('CODIGO MUPI', ''),
                    'descripcion': datos.get('descripción', ''),
                }
                pines[pin] = d
            else:
                capa2 = capa
                if tag == 'Folder':
                    capa2 = nombre_de(hijo) or capa
                walk(hijo, capa2)

    walk(raiz, '')
    return pines


def _pines_desde_xlsx(ruta):
    wb = load_workbook(ruta, read_only=True, data_only=True)
    ws = wb['Pines'] if 'Pines' in wb.sheetnames else wb[wb.sheetnames[0]]
    enc = [str(ws.cell(1, c).value or '').strip().upper() for c in range(1, ws.max_column + 1)]
    pines = {}
    for r in range(2, ws.max_row + 1):
        fila = {}
        for c, nom in enumerate(enc, 1):
            fila[nom] = ws.cell(r, c).value
        pin = str(fila.get('PIN') or '').strip()
        if pin:
            pines[pin] = {k: ('' if v is None else str(v)).strip() if not isinstance(v, (int, float)) else v
                          for k, v in fila.items()}
    wb.close()
    return pines


def _mid_desde_url(url):
    m = re.search(r'mid=([A-Za-z0-9_\-]+)', url or '')
    return m.group(1) if m else None


def cargar_pines(ruta_pines):
    """Carga pines desde .xlsx | .kml | .txt (con URL del mapa). Devuelve dict PIN->datos."""
    if not ruta_pines or not os.path.exists(ruta_pines):
        raise RuntimeError('No existe la ruta de GoogleMapsPines: {}'.format(ruta_pines))
    ext = os.path.splitext(ruta_pines)[1].lower()
    if ext == '.xlsx':
        pines = _pines_desde_xlsx(ruta_pines)
    elif ext == '.kml':
        pines = _pines_desde_kml(ruta_pines)
    else:
        texto = U.leer_texto_auto(ruta_pines)
        url = re.search(r'https?://\S+', texto)
        if not url:
            raise RuntimeError('El archivo de pines no contiene una URL de Google Maps.')
        mid = _mid_desde_url(url.group(0))
        if not mid:
            raise RuntimeError('No se pudo extraer el id (mid) del mapa.')
        if en_mock():
            raise RuntimeError('MODO_MOCK: usa un .kml o .xlsx local para pines.')
        kml_url = 'https://www.google.com/maps/d/kml?mid={}&forcekml=1'.format(mid)
        req = urllib.request.Request(kml_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=90) as r:
            datos = r.read()
        tmp = os.path.join(os.environ.get('TEMP', '.'), 'cvlo_pines_{}.kml'.format(mid))
        with open(tmp, 'wb') as f:
            f.write(datos)
        pines = _pines_desde_kml(tmp)
    limpios = {}
    for pin, d in pines.items():
        try:
            d['lat'] = float(str(d.get('LATITUD', '')).replace(',', '.'))
            d['lon'] = float(str(d.get('LONGITUD', '')).replace(',', '.'))
        except Exception:
            continue
        limpios[pin] = d
    if not limpios:
        raise RuntimeError('No se cargaron pines con coordenadas válidas.')
    return limpios


def cargar_kml_datos(ruta_kml):
    if not ruta_kml or not os.path.exists(ruta_kml):
        return {}
    try:
        return _pines_desde_kml(ruta_kml)
    except Exception:
        return {}


MAPA_TIPO = {'PARABUS': 'PARABUSES', 'BILLBOARD': 'PANORAMICOS', 'ISA MUPI': 'MUPIS', 'MUPI': 'MUPIS',
             'CARTELERA': 'CARTELERAS', 'PANTALLA DIGITAL': 'PANTALLAS', 'KIOSKO': 'KIOSKOS', 'PUENTE': 'PUENTES',
             'VALLA': 'PANORAMICOS', 'ESPECTACULAR': 'PANORAMICOS', 'PANORAMICO': 'PANORAMICOS',
             'PANORÁMICO': 'PANORAMICOS', 'MURO': 'MUROS', 'PARABÚS': 'PARABUSES', 'PARABUSES': 'PARABUSES'}


def marca_de(pin, d):
    if pin.startswith('DF_'):
        return 'DIDI'
    if pin.startswith('PF_COP'):
        return (d.get('CLIENTE') or 'COPPEL').strip()
    if pin.startswith(('COM', 'Punto')):
        return (d.get('MARCA') or 'COMPETENCIA').strip()
    return 'OTROS'


def tipo_de(pin, d):
    m = str(d.get('MEDIO') or '').upper().strip()
    if m in MAPA_TIPO:
        return MAPA_TIPO[m]
    if pin.startswith('PF_COP'):
        return 'PARABUSES'
    return 'OTROS' if not m else U.sanear_carpeta(m.replace(' ', ''))


def grupo_de(pin, d):
    return U.sanear_carpeta('{} {} MTY'.format(marca_de(pin, d), tipo_de(pin, d)))


# ---------------------------------------------------------------- idxfotos
IDX_HEADERS = ['NombreFoto', 'Descripción de lo que anuncia', 'Dirección', 'Tipo de Anuncio', 'Latitud',
               'Longitud', 'Referencias visuales o notas', 'Verificado', 'PIN']


def idx_abrir_o_crear(ruta):
    if ruta and os.path.exists(ruta):
        wb = load_workbook(ruta)
        ws = wb['idxfotos'] if 'idxfotos' in wb.sheetnames else wb[wb.sheetnames[0]]
        return wb, ws
    wb = Workbook()
    ws = wb.active
    ws.title = 'idxfotos'
    ws.append(IDX_HEADERS)
    fuente = Font(bold=True, color='FFFFFF', size=11)
    relleno = PatternFill('solid', fgColor='C00000')
    for c in range(1, len(IDX_HEADERS) + 1):
        celda = ws.cell(1, c)
        celda.font = fuente
        celda.fill = relleno
        celda.alignment = Alignment(horizontal='center', vertical='center')
    return wb, ws


def idx_filas(ws):
    """dict nombre -> numero de fila, y set de verificados."""
    pos = {}
    verif = set()
    for r in range(2, ws.max_row + 1):
        nombre = str(ws.cell(r, 1).value or '').strip()
        if not nombre:
            continue
        pos[nombre] = r
        v = ws.cell(r, 8).value
        if v and str(v).strip():
            verif.add(nombre)
    return pos, verif


def idx_upsert(ws, pos, fila):
    nombre = str(fila[0])
    if nombre in pos:
        r = pos[nombre]
        for c, val in enumerate(fila, 1):
            ws.cell(r, c).value = val
    else:
        ws.append(fila)
        pos[nombre] = ws.max_row


def idx_guardar(wb, ruta):
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    try:
        wb.save(ruta)
        return ruta
    except PermissionError:
        time.sleep(3)
        try:
            wb.save(ruta)
            return ruta
        except PermissionError:
            alt = ruta.replace('.xlsx', '_nuevo.xlsx')
            wb.save(alt)
            return alt
