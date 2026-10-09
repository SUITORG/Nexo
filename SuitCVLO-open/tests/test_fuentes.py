# -*- coding: utf-8 -*-
"""Botón 'Ver info de las fuentes': resumen del KML y los dos endpoints."""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

os.environ['CVLO_MOCK'] = '1'

from PIL import Image  # noqa: E402
import motor  # noqa: E402
from tests.test_flujo_mock import crear_fuente, KML_DE_PRUEBA  # noqa: E402


class TestKmlResumen(unittest.TestCase):
    def test_resumen_del_kml(self):
        base = tempfile.mkdtemp(prefix='cvlo_kmlres_')
        ruta = os.path.join(base, 'p.kml')
        with open(ruta, 'w', encoding='utf-8') as f:
            f.write(KML_DE_PRUEBA)
        r = motor.kml_resumen(motor.cargar_kml_datos(ruta))
        self.assertEqual(r['total'], 2)
        self.assertEqual(r['capas'], {'PRUEBA (1)': 2})   # sin '.xlsx'
        self.assertEqual(r['cara'], 1)
        self.assertEqual(r['sin_lat'], 0)
        # lo que aporta el KML (identidad/geometría fuera)
        self.assertEqual(r['campos'].get('CARA'), 1)
        self.assertEqual(r['campos'].get('MEDIO'), 2)
        self.assertNotIn('PIN', r['campos'])
        self.assertNotIn('LATITUD', r['campos'])

    def test_resumen_kml_vacio(self):
        r = motor.kml_resumen({})
        self.assertEqual(r, {'total': 0, 'capas': {}, 'cara': 0, 'sin_lat': 0, 'campos': {}})


class TestApiFuentes(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import app as appmod
        appmod.app.config['TESTING'] = True
        cls.cliente = appmod.app.test_client()

    def _base(self, n=2, con_kml=True):
        base = tempfile.mkdtemp(prefix='cvlo_fuentes_')
        fotos, pines, kml = crear_fuente(base, [
            ('IMG_7401.jpg', '2026-09-14', 25.70002, -100.3500),
            ('IMG_7402.jpg', '2026-09-14', 25.70006, -100.3500)][:n])
        if con_kml:
            with open(kml, 'w', encoding='utf-8') as f:
                f.write(KML_DE_PRUEBA)
        return base, fotos, pines, kml

    def test_fuentes_rapidas(self):
        base, fotos, pines, kml = self._base()
        r = self.cliente.get('/api/fuentes', query_string={
            'ruta': fotos, 'pines': pines, 'kml': kml,
            'idx': os.path.join(base, 'no-existe.xlsx')})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        d = r.get_json()
        self.assertEqual(d['fotos'], 2)
        self.assertEqual(d['pines'], 2)
        self.assertEqual(d['kml']['total'], 2)
        self.assertEqual(d['kml']['cara'], 1)
        self.assertEqual(d['indice'], {'filas': 0, 'verificadas': 0, 'no_gps': 0, 'duplicadas': 0})

    def test_indice_detecta_fotos_duplicadas(self):
        """IMG_0139.jpg y img_0139.jpg son la misma foto: cuenta como duplicada (info revuelta)."""
        import os
        import tempfile
        from openpyxl import Workbook
        import app as appmod
        d = tempfile.mkdtemp(prefix='cvlo_dup_')
        ruta = os.path.join(d, 'idx.xlsx')
        wb = Workbook()
        ws = wb.active
        ws.append(['NombreFoto', 'b', 'c', 'd', 'e', 'f', 'g', 'Verificado', 'PIN'])
        ws.append(['IMG_0139.jpg', '', '', '', '', '', '', '', 'P1'])
        ws.append(['img_0139.jpg', '', '', '', '', '', '', '✓', 'P1'])
        ws.append(['otra.jpg', '', '', '', '', '', '', '✓', 'P1'])
        wb.save(ruta)
        filas, verif, no_gps, _stems, dup, filas_idx = appmod._leer_idx(ruta)
        self.assertEqual((filas, verif, dup), (3, 2, 1))
        self.assertEqual(filas_idx['img_0139']['pin'], 'P1')     # con PIN y verificado
        self.assertTrue(filas_idx['img_0139']['verif'])
        self.assertEqual(filas_idx['otra']['pin'], 'P1')
        # y sale en /api/fuentes
        r = self.cliente.get('/api/fuentes', query_string={'idx': ruta})
        self.assertEqual(r.get_json()['indice']['duplicadas'], 1)

    def test_fuentes_ruta_inexistente_devuelve_400(self):
        r = self.cliente.get('/api/fuentes', query_string={'ruta': r'C:\no\existe\xyz'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('no existe', r.get_json()['error'])

    def test_fuentes_gps_desglosa(self):
        """2 con GPS + 1 sin GPS -> con_gps=2, sin_gps=1, fuera_radio=0."""
        base, fotos, pines, kml = self._base(n=2)
        Image.new('RGB', (320, 240), (240, 240, 240)).save(os.path.join(fotos, 'IMG_7403.jpg'))
        r = self.cliente.get('/api/fuentes/gps', query_string={
            'ruta': fotos, 'pines': pines, 'radio': 250,
            'salida': os.path.join(base, 'salida'),
            'idx': os.path.join(base, 'no-existe.xlsx')})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        d = r.get_json()
        self.assertEqual(d['con_gps'], 2)
        self.assertEqual(d['sin_gps'], 1)
        self.assertEqual(d['fuera_radio'], 0)
        self.assertEqual(d['radio'], 250)


    def test_campanas_salen_del_kml(self):
        """Con KML, las campañas salen de ahí (no se mezclan con el mapa de pines)."""
        from openpyxl import Workbook
        base, fotos, pines, kml = self._base()
        idx = os.path.join(base, 'idx.xlsx')
        wb = Workbook()
        ws = wb.active
        ws.append(['NombreFoto', 'b', 'c', 'd', 'Latitud', 'Longitud', 'g', 'Verificado', 'PIN'])
        ws.append(['IMG_7401.jpg', '', '', '', 25.70002, -100.35, '', '✓', 'DF_09_TEST_1'])
        wb.save(idx)
        r = self.cliente.get('/api/fuentes', query_string={'pines': pines, 'kml': kml, 'idx': idx})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        camp = r.get_json()['campanas']
        # solo la campaña del KML (DF_09_TEST_1 es de los pines: no cuenta aquí)
        self.assertEqual([c['nombre'] for c in camp], ['PRUEBA (1)'])
        self.assertEqual(camp[0]['pines'], 2)
        self.assertEqual(camp[0]['auditados'], 0)

    def test_campanas_sin_kml_usan_el_mapa_de_pines(self):
        from openpyxl import Workbook
        base, fotos, pines, kml = self._base(con_kml=False)
        idx = os.path.join(base, 'idx.xlsx')
        wb = Workbook()
        ws = wb.active
        ws.append(['NombreFoto', 'b', 'c', 'd', 'Latitud', 'Longitud', 'g', 'Verificado', 'PIN'])
        ws.append(['IMG_7401.jpg', '', '', '', 25.70002, -100.35, '', '✓', 'DF_09_TEST_1'])
        wb.save(idx)
        r = self.cliente.get('/api/fuentes', query_string={'pines': pines, 'idx': idx})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        camp = {c['nombre']: c for c in r.get_json()['campanas']}
        self.assertEqual(list(camp), ['DIDI PRUEBA'])
        self.assertEqual(camp['DIDI PRUEBA']['pines'], 2)
        self.assertEqual(camp['DIDI PRUEBA']['auditados'], 1)      # 1 pin ya auditado
        self.assertEqual(camp['DIDI PRUEBA']['verificadas'], 1)

    def test_gps_por_campana_da_distancia_y_pendientes(self):
        """La fase 2 devuelve la distancia media foto<->pin y las pendientes por campaña."""
        base, fotos, pines, kml = self._base(n=2)
        r = self.cliente.get('/api/fuentes/gps', query_string={
            'ruta': fotos, 'pines': pines, 'radio': 250,
            'salida': os.path.join(base, 'salida'),
            'idx': os.path.join(base, 'no-existe.xlsx')})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        c = r.get_json()['campanas']
        self.assertIn('DIDI PRUEBA', c)
        self.assertEqual(c['DIDI PRUEBA']['n'], 2)              # las 2 fotos cruzan con DF_09_TEST_1
        self.assertEqual(c['DIDI PRUEBA']['pendientes'], 2)     # índice vacío -> ninguna verificada
        self.assertIsInstance(c['DIDI PRUEBA']['dist_prom'], float)
        self.assertLess(c['DIDI PRUEBA']['dist_prom'], 20)      # a unos metros del pin
        self.assertNotIn('PRUEBA (1)', c)                       # esas fotos no caen en el KML


    def test_aviso_si_pines_y_kml_son_mapas_distintos(self):
        """Si el mapa de pines y el KML no comparten PINs, el panel avisa (cruce != campañas)."""
        base, fotos, pines, kml = self._base()
        r = self.cliente.get('/api/fuentes', query_string={'pines': pines, 'kml': kml})
        d = r.get_json()
        self.assertIn('aviso_mapa', d)
        self.assertIn('NO están en el KML', d['aviso_mapa'])
        # sin KML no hay de qué comparar
        r2 = self.cliente.get('/api/fuentes', query_string={'pines': pines})
        self.assertNotIn('aviso_mapa', r2.get_json())

    def test_hechas_y_pendientes_para_el_avance(self):
        """AVANCE = hechas (copiadas o leídas) ÷ total: 100% cuando no queda nada por procesar."""
        import shutil
        from unittest import mock
        base, fotos, pines, kml = self._base(n=2)
        salida = os.path.join(base, 'salida')
        destino = os.path.join(salida, '260901', 'GRUPO', 'DF_1')
        os.makedirs(destino, exist_ok=True)
        shutil.copy2(os.path.join(fotos, 'IMG_7401.jpg'), os.path.join(destino, 'IMG_7401.jpg'))
        import app as appmod
        with mock.patch.object(appmod, 'TRABAJO', tempfile.mkdtemp(prefix='cvlo_trab_')):
            r = self.cliente.get('/api/fuentes', query_string={
                'ruta': fotos, 'salida': salida,
                'idx': os.path.join(base, 'no-existe.xlsx')})
        d = r.get_json()
        self.assertEqual(d['fotos'], 2)
        self.assertEqual(d['hechas'], 1)          # la copiada
        self.assertEqual(d['pendientes'], 1)      # la que falta
        self.assertEqual(d['hechas'] + d['pendientes'], d['fotos'])   # cuadra

    def test_campanas_leidos_para_el_porcentaje(self):
        """% avance = pines con >=1 foto LEÍDA (aunque no esté verificada)."""
        from openpyxl import Workbook
        base, fotos, pines, kml = self._base()
        idx = os.path.join(base, 'idx.xlsx')
        wb = Workbook()
        ws = wb.active
        ws.append(['NombreFoto', 'b', 'c', 'd', 'Latitud', 'Longitud', 'g', 'Verificado', 'PIN'])
        ws.append(['IMG_7401.jpg', '', '', '', 25.7, -100.35, '', '', 'DF_KML_A'])  # sin ✓
        wb.save(idx)
        r = self.cliente.get('/api/fuentes', query_string={'kml': kml, 'idx': idx})
        camp = {c['nombre']: c for c in r.get_json()['campanas']}['PRUEBA (1)']
        self.assertEqual(camp['pines'], 2)
        self.assertEqual(camp['leidos'], 1)        # cuenta para el %
        self.assertEqual(camp['auditados'], 0)     # pero no está verificada


class TestPdfInforme(unittest.TestCase):
    """POST /api/fuentes/pdf -> analisis_fuentes_AAAA-MM-DD_HH-MM.pdf en la Ruta de Salida."""

    @classmethod
    def setUpClass(cls):
        os.environ.setdefault('CVLO_TRABAJO', tempfile.mkdtemp(prefix='cvlo_app_trabajo_'))
        import app as appmod
        appmod.app.config['TESTING'] = True
        cls.cliente = appmod.app.test_client()

    def test_pdf_se_guarda_en_la_salida_con_fecha(self):
        import re
        salida = tempfile.mkdtemp(prefix='cvlo_pdfout_')
        r = self.cliente.post('/api/fuentes/pdf', json={
            'salida': salida,
            'lineas': ['RUTA      C:\\fotos', 'AVANCE 27% · hechas 463 de 1696',
                       'CAMPAÑAS', '  COPPEL PARABUSES MTY (1)   159', 'con ñ, á y ✓ y ≥']})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        d = r.get_json()
        self.assertRegex(d['archivo'], r'^analisis_fuentes_\d{4}-\d{2}-\d{2}_\d{2}-\d{2}\.pdf$')
        self.assertEqual(os.path.dirname(d['ruta']), salida)
        self.assertTrue(os.path.exists(d['ruta']))
        b = open(d['ruta'], 'rb').read()
        self.assertEqual(b[:8], b'%PDF-1.4')
        self.assertTrue(b.rstrip().endswith(b'%%EOF'))
        self.assertIn(b'Generado: ', b)                       # fecha y hora dentro
        self.assertIn(b'hechas 463 de 1696', b)               # contenido del análisis
        # el xref apunta a la tabla y cada offset lleva su "N 0 obj"
        xref = int(b.split(b'startxref')[1].split(b'%%EOF')[0].strip())
        self.assertEqual(b[xref:xref + 4], b'xref')
        for off in re.findall(rb'^(\d{10}) 00000 n', b[xref:], re.M):
            i = int(off)
            self.assertTrue(b[i:i + 1].isdigit(), 'offset roto: %r' % off)
        self.assertGreater(d['bytes'], 900)

    def test_pdf_valida_la_entrada(self):
        self.assertEqual(self.cliente.post('/api/fuentes/pdf', json={'lineas': ['x']}).status_code, 400)
        self.assertEqual(
            self.cliente.post('/api/fuentes/pdf',
                              json={'salida': r'C:\no\existe\x', 'lineas': ['x']}).status_code, 400)
        salida = tempfile.mkdtemp(prefix='cvlo_pdfout2_')
        self.assertEqual(
            self.cliente.post('/api/fuentes/pdf',
                              json={'salida': salida, 'lineas': []}).status_code, 400)


if __name__ == '__main__':
    unittest.main()
