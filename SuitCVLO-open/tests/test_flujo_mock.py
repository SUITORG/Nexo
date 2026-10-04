# -*- coding: utf-8 -*-
"""E2E en modo MOCK: proceso completo + API de la app."""
import json
import os
import sys
import tempfile
import time
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

os.environ['CVLO_MOCK'] = '1'

from PIL import Image
from openpyxl import load_workbook

import motor
import proceso


def crear_fuente(base, fotos_datos):
    """fotos_datos: [(nombre, fecha, lat, lon)] + mock_gps.json + pines.xlsx."""
    fotos = os.path.join(base, 'fotos')
    os.makedirs(fotos, exist_ok=True)
    mock = {}
    for i, (nombre, fecha, lat, lon) in enumerate(fotos_datos):
        Image.new('RGB', (320, 240), (240 - i * 10, 240, 240)).save(os.path.join(fotos, nombre), quality=90)
        mock[nombre] = {'lat': lat, 'lon': lon, 'fecha': fecha}
    with open(os.path.join(fotos, 'mock_gps.json'), 'w', encoding='utf-8') as f:
        json.dump(mock, f)
    # pines
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Pines'
    ws.append(['PIN', 'LATITUD', 'LONGITUD', 'MEDIO', 'FORMATO', 'ID', 'UBICACIÓN', 'CIUDAD', 'CAPA'])
    ws.append(['DF_09_TEST_1', 25.7000, -100.3500, 'PARABUS', 'FIJO', 'T-1', 'CALLE UNO', 'MTY', 'DIDI PRUEBA'])
    ws.append(['DF_09_TEST_2', 25.7100, -100.3600, 'BILLBOARD', 'FIJO', 'T-2', 'CALLE DOS', 'MTY', 'DIDI PRUEBA'])
    pines = os.path.join(base, 'pines.xlsx')
    wb.save(pines)
    kml = os.path.join(base, 'ruta.kml')
    open(kml, 'w', encoding='utf-8').write('<kml/>')
    return fotos, pines, kml


class TestFlujoMock(unittest.TestCase):

    def test_corrida_completa(self):
        base = tempfile.mkdtemp(prefix='cvlo_e2e_')
        fotos, pines, kml = crear_fuente(base, [
            ('IMG_0001.jpg', '2026-09-14', 25.70002, -100.3500),
            ('IMG_0002.jpg', '2026-09-14', 25.70006, -100.3500),
            ('IMG_0003.jpg', '2026-09-20', 25.71002, -100.3600)])
        salida = os.path.join(base, 'salida')
        idx = os.path.join(base, 'idxfotos.xlsx')
        trabajo = os.path.join(base, 'trabajo')

        params = {'modo': 'local', 'ruta_local': fotos, 'salida': salida, 'pines': pines,
                  'drive_pics': '', 'kml': kml, 'idx': idx, 'n_fotos': 10, 'radio': 250,
                  'cap': 10, 'pausa': 0, 'trabajo': trabajo,
                  '_mock_gps': json.load(open(os.path.join(fotos, 'mock_gps.json'), encoding='utf-8'))}
        eventos = []
        resumen = proceso.ejecutar_corrida(params, lambda c: eventos.append(c), lambda m: eventos.append(m))

        self.assertEqual(resumen['leidas'], 3)
        self.assertEqual(resumen['copiadas'], 3)
        # estructura 260901 y 260902 (según fecha de foto)
        f1 = os.path.join(salida, '260901', 'auditoria', 'DIDI PARABUSES MTY', 'DF_09_TEST_1')
        f3 = os.path.join(salida, '260902', 'auditoria', 'DIDI PANORAMICOS MTY', 'DF_09_TEST_2')
        self.assertTrue(os.path.exists(os.path.join(f1, 'IMG_0001.jpg')))
        self.assertTrue(os.path.exists(os.path.join(f1, 'IMG_0002.jpg')))
        self.assertTrue(os.path.exists(os.path.join(f3, 'IMG_0003.jpg')))
        self.assertTrue(os.path.exists(os.path.join(f1, 'info.txt')))
        info = open(os.path.join(f1, 'info.txt'), encoding='utf-8').read()
        self.assertIn('PIN: DF_09_TEST_1', info)
        self.assertIn('CAMPAÑA:', info)
        # idxfotos creado con 3 filas verificadas
        wb = load_workbook(idx)
        ws = wb['idxfotos']
        self.assertEqual(ws.max_row - 1, 3)
        verif = [ws.cell(r, 8).value for r in range(2, ws.max_row + 1)]
        self.assertEqual(sum(1 for v in verif if v), 3)

        # segunda corrida: no repite nada
        params2 = dict(params)
        params2['_mock_gps'] = params['_mock_gps']
        resumen2 = proceso.ejecutar_corrida(params2, lambda c: None, lambda m: None)
        self.assertEqual(resumen2['leidas'], 0)

    def test_tope_por_pin(self):
        base = tempfile.mkdtemp(prefix='cvlo_cap_')
        fotos, pines, kml = crear_fuente(base, [
            ('IMG_1001.jpg', '2026-09-14', 25.70002, -100.3500),
            ('IMG_1002.jpg', '2026-09-14', 25.70004, -100.3500),
            ('IMG_1003.jpg', '2026-09-14', 25.70006, -100.3500)])
        salida = os.path.join(base, 'salida')
        idx = os.path.join(base, 'idxfotos.xlsx')
        params = {'modo': 'local', 'ruta_local': fotos, 'salida': salida, 'pines': pines,
                  'drive_pics': '', 'kml': kml, 'idx': idx, 'n_fotos': 10, 'radio': 250,
                  'cap': 2, 'pausa': 0, 'trabajo': os.path.join(base, 'trabajo'),
                  '_mock_gps': json.load(open(os.path.join(fotos, 'mock_gps.json'), encoding='utf-8'))}
        resumen = proceso.ejecutar_corrida(params, lambda c: None, lambda m: None)
        self.assertEqual(resumen['leidas'], 3)
        self.assertEqual(resumen['copiadas'], 2)
        self.assertEqual(resumen['excedidas'], 1)
        wb = load_workbook(idx)
        ws = wb['idxfotos']
        blancos = [r for r in range(2, ws.max_row + 1) if not ws.cell(r, 8).value]
        self.assertEqual(len(blancos), 1)
        nota = str(ws.cell(blancos[0], 7).value or '')
        self.assertIn('excede', nota)


class TestApiMock(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        os.environ['CVLO_MOCK'] = '1'
        os.environ['CVLO_TRABAJO'] = tempfile.mkdtemp(prefix='cvlo_app_trabajo_')
        import app as appmod
        cls.appmod = appmod
        appmod.app.config['TESTING'] = True
        cls.cliente = appmod.app.test_client()

    def test_api_local(self):
        base = tempfile.mkdtemp(prefix='cvlo_api_')
        fotos, pines, kml = crear_fuente(base, [('IMG_2001.jpg', '2026-09-14', 25.70002, -100.3500)])
        salida = os.path.join(base, 'salida')
        idx = os.path.join(base, 'idxfotos.xlsx')
        r = self.cliente.post('/api/corridas', data={
            'modo': 'local', 'ruta_local': fotos, 'salida': salida, 'pines': pines,
            'kml': kml, 'idx': idx, 'n_fotos': '5'}, content_type='multipart/form-data')
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True)[:200])
        jid = r.get_json()['id']
        estado = None
        tope = time.time() + 60
        while time.time() < tope:
            estado = self.cliente.get('/api/corridas/' + jid).get_json()
            if estado['estado'] in ('listo', 'error'):
                break
            time.sleep(0.3)
        self.assertEqual(estado['estado'], 'listo', estado.get('error'))
        self.assertEqual(estado['completadas'], 1)
        self.assertTrue(estado['resumen'] is not None)
        self.assertEqual(estado['resumen']['leidas'], 1)
        # salud
        self.assertEqual(self.cliente.get('/api/estado').status_code, 200)
        self.assertEqual(self.cliente.get('/').status_code, 200)


class TestComentarioYFechaSubida(unittest.TestCase):
    """Lo pedido por el usuario: COMENTARIO en info.txt y FechaHoraSubida en idxfotos.xlsx."""

    def _params(self, base, idx):
        fotos, pines, kml = crear_fuente(base, [('IMG_6001.jpg', '2026-09-14', 25.70002, -100.3500)])
        return {'modo': 'local', 'ruta_local': fotos, 'salida': os.path.join(base, 'salida'),
                'pines': pines, 'drive_pics': '', 'kml': kml, 'idx': idx,
                'n_fotos': 10, 'radio': 250, 'cap': 10, 'pausa': 0,
                'trabajo': os.path.join(base, 'trabajo'),
                '_mock_gps': json.load(open(os.path.join(fotos, 'mock_gps.json'), encoding='utf-8'))}

    def test_info_txt_lleva_comentario(self):
        base = tempfile.mkdtemp(prefix='cvlo_coment_')
        idx = os.path.join(base, 'idxfotos.xlsx')
        proceso.ejecutar_corrida(self._params(base, idx), lambda c: None, lambda m: None)
        ruta = os.path.join(base, 'salida', '260901', 'auditoria',
                            'DIDI PARABUSES MTY', 'DF_09_TEST_1', 'info.txt')
        self.assertTrue(os.path.exists(ruta), ruta)
        texto = open(ruta, encoding='utf-8').read()
        self.assertIn('COMENTARIO:', texto)
        # el comentario es lo que describió la IA que promociona (modo mock)
        self.assertIn('promoción simulada', texto)
        # y va DENTRO del bloque de la foto, antes de la sección final COMENTARIOS:
        self.assertLess(texto.index('COMENTARIO:'), texto.index('COMENTARIOS:'))

    def test_info_txt_lleva_foto(self):
        """Cada bloque empieza con FOTO: <nombre> — tantas como fotos en la carpeta."""
        base = tempfile.mkdtemp(prefix='cvlo_foto_')
        idx = os.path.join(base, 'idxfotos.xlsx')
        proceso.ejecutar_corrida(self._params(base, idx), lambda c: None, lambda m: None)
        carpeta = os.path.join(base, 'salida', '260901', 'auditoria',
                               'DIDI PARABUSES MTY', 'DF_09_TEST_1')
        texto = open(os.path.join(carpeta, 'info.txt'), encoding='utf-8').read()
        fotos = [f for f in os.listdir(carpeta) if f.lower().endswith('.jpg')]
        self.assertEqual(len(fotos), 1)
        self.assertIn('FOTO: %s' % fotos[0], texto)
        # el bloque arranca con FOTO, antes que CAMPAÑA
        bloque = texto.split('-' * 40)[1]
        self.assertTrue(bloque.strip().startswith('FOTO:'), bloque[:80])

    def test_info_txt_sin_comentario_si_switch_desactivado(self):
        """Switch de Opciones avanzadas en OFF: no sale COMENTARIO pero sí FOTO."""
        base = tempfile.mkdtemp(prefix='cvlo_swtch_')
        params = self._params(base, os.path.join(base, 'idxfotos.xlsx'))
        params['comentario'] = False
        proceso.ejecutar_corrida(params, lambda c: None, lambda m: None)
        carpeta = os.path.join(base, 'salida', '260901', 'auditoria',
                               'DIDI PARABUSES MTY', 'DF_09_TEST_1')
        texto = open(os.path.join(carpeta, 'info.txt'), encoding='utf-8').read()
        self.assertIn('FOTO:', texto)
        self.assertNotIn('COMENTARIO:', texto)
        self.assertIn('CAMPAÑA:', texto)

    def test_rellena_fecha_hora_subida_sin_crear_columna_nueva(self):
        from openpyxl import Workbook
        base = tempfile.mkdtemp(prefix='cvlo_subida_')
        idx = os.path.join(base, 'idxfotos.xlsx')
        wb = Workbook()
        ws = wb.active
        ws.title = 'Sheet1'                     # como el xlsx real del usuario
        ws.append(['NombreFoto', 'Descripción de lo que anuncia', 'Dirección', 'Tipo de Anuncio',
                   'Latitud', 'Longitud', 'Referencias visuales o notas', 'Verificado', 'PIN',
                   'FechaHoraFoto', 'FechaHoraSubida'])
        ws.append(['vieja.jpg', 'que anuncia', 'calle', 'mupi', 1, 2, 'refs', 'ok',
                   'DF_09_TEST_1', '2026-09-01 08:00', '2026-09-01 09:00'])
        wb.save(idx)

        proceso.ejecutar_corrida(self._params(base, idx), lambda c: None, lambda m: None)

        wb2 = load_workbook(idx)
        ws2 = wb2[wb2.sheetnames[0]]
        # única columna nueva = Marca (al final); la fecha de subida NO creó la suya
        self.assertEqual(ws2.max_column, 12)
        self.assertEqual(ws2.cell(1, 10).value, 'FechaHoraFoto')
        self.assertEqual(ws2.cell(1, 11).value, 'FechaHoraSubida')
        self.assertEqual(ws2.cell(1, 12).value, 'Marca')
        self.assertEqual(ws2.cell(2, 11).value, '2026-09-01 09:00')   # su dato, intacto
        self.assertTrue(str(ws2.cell(2, 12).value or '') != '2026-09-01 09:00')  # no es una fecha
        nueva = ws2.max_row
        self.assertRegex(str(ws2.cell(nueva, 11).value or ''),
                         r'^\d{4}-\d{2}-\d{2} \d{2}:\d{2}$')          # hora de subida puesta
        self.assertTrue(ws2.cell(nueva, 10).value in (None, ''))       # FechaHoraFoto no se inventa
        wb2.close()


class TestMarcaYCampana(unittest.TestCase):
    """PROVEEDOR = el anunciante/marca; CAMPAÑA = el eslogan (prompt de 8 líneas)."""

    def test_formato_nuevo_separa_marca_y_campana(self):
        texto = ('ANUNCIO: SI\n'
                 'MARCA: Farmacias del Ahorro\n'
                 'CAMPAÑA: Siempre cerca de ti\n'
                 'QUE ANUNCIA: promociona servicios\n'
                 'TEXTO VISIBLE: Siempre cerca de ti\n'
                 'REFERENCIAS UBICACION: esquina\n'
                 'TIPO SOPORTE: parabús\nCARA: frente')
        bloque = proceso._bloque_info(texto, {'PIN': 'DF_09_TEST_1'}, {}, 25.7, -100.35,
                                      foto='IMG_1.jpg')
        self.assertTrue(bloque.startswith('FOTO: IMG_1.jpg'))
        self.assertIn('CAMPAÑA: Siempre cerca de ti', bloque)
        self.assertIn('PROVEEDOR: Farmacias del Ahorro', bloque)

    def test_lectura_legada_usa_la_marca_para_los_dos(self):
        """Las lecturas viejas traen 'MARCA CAMPAÑA:' fusionado (una sola etiqueta)."""
        texto = ('ANUNCIO: SI\nMARCA CAMPAÑA: Coppel\nQUE ANUNCIA: promo\n'
                 'TEXTO VISIBLE: Lee\nREFERENCIAS UBICACION: n/a\n'
                 'TIPO SOPORTE: mupi\nCARA: frente')
        bloque = proceso._bloque_info(texto, {'PIN': 'PF_COP_1', 'CLIENTE': ''}, {},
                                      19.4, -99.1, foto='a.jpg')
        # 'MARCA' casa por prefijo con 'MARCA CAMPAÑA' -> PROVEEDOR se llena igual
        self.assertIn('PROVEEDOR: Coppel', bloque)
        # sin campo CAMPAÑA explicito -> fallback a la marca (como hasta ahora)
        self.assertIn('CAMPAÑA: Coppel', bloque)

    def test_proveedor_sin_ia_cae_en_la_regla_del_pin(self):
        texto = ('ANUNCIO: SI\nMARCA: no visible\nCAMPAÑA: no visible\n'
                 'QUE ANUNCIA: n/a\nTEXTO VISIBLE: n/a\nREFERENCIAS UBICACION: n/a\n'
                 'TIPO SOPORTE: otro\nCARA: frente')
        bloque = proceso._bloque_info(texto, {'PIN': 'DF_09_TEST_1'}, {}, 25.7, -100.35)
        # DF_* -> DIDI por la regla de motor.marca_de
        self.assertIn('PROVEEDOR: DIDI', bloque)
        # con CAMPAÑA explicita pero 'no visible', se respeta el valor de la IA
        self.assertIn('CAMPAÑA: no visible', bloque)


KML_DE_PRUEBA = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    '<kml xmlns="http://www.opengis.net/kml/2.2"><Document>\n'
    '  <Folder><name>PRUEBA (1).xlsx</name>\n'
    '    <Placemark><name>DF_KML_A</name><ExtendedData>\n'
    '      <Data name="LATITUD"><value>25.7000</value></Data>\n'
    '      <Data name="LONGITUD"><value>-100.3500</value></Data>\n'
    '      <Data name="CARA"><value>Cara A</value></Data>\n'
    '      <Data name="MEDIO"><value>PARABUS</value></Data>\n'
    '    </ExtendedData></Placemark>\n'
    '    <Placemark><name>DF_KML_B</name><ExtendedData>\n'
    '      <Data name="LATITUD"><value>25.7100</value></Data>\n'
    '      <Data name="LONGITUD"><value>-100.3600</value></Data>\n'
    '      <Data name="MEDIO"><value>MUPI</value></Data>\n'
    '    </ExtendedData></Placemark>\n'
    '  </Folder>\n'
    '</Document></kml>\n')


class TestStatsYNoGps(unittest.TestCase):
    """Pantalla: nº de fotos en la ruta / con GPS / sin GPS / fuera de radio / PINs.
    Índice: las fotos sin GPS se registran con 'No GPS' en la columna Verificado."""

    def _montar(self):
        base = tempfile.mkdtemp(prefix='cvlo_stats_')
        fotos, pines, kml = crear_fuente(base, [
            ('IMG_7101.jpg', '2026-09-14', 25.70002, -100.3500),
            ('IMG_7102.jpg', '2026-09-14', 25.70006, -100.3500)])
        # tercera foto SIN gps: no está en mock_gps.json
        Image.new('RGB', (320, 240), (240, 240, 240)).save(os.path.join(fotos, 'IMG_7103.jpg'))
        params = {'modo': 'local', 'ruta_local': fotos, 'salida': os.path.join(base, 'salida'),
                  'pines': pines, 'drive_pics': '', 'kml': kml,
                  'idx': os.path.join(base, 'idxfotos.xlsx'), 'n_fotos': 10, 'radio': 250,
                  'cap': 10, 'pausa': 0, 'trabajo': os.path.join(base, 'trabajo'),
                  '_mock_gps': json.load(open(os.path.join(fotos, 'mock_gps.json'), encoding='utf-8'))}
        return base, params

    @staticmethod
    def _stats(eventos):
        s = [e['stats'] for e in eventos if isinstance(e, dict) and 'stats' in e]
        return s[-1] if s else None

    def test_stats_se_reportan(self):
        _base, params = self._montar()
        eventos = []
        proceso.ejecutar_corrida(params, eventos.append, lambda m: None)
        s = self._stats(eventos)
        self.assertIsNotNone(s, 'no se emitió progreso stats')
        self.assertEqual(s['fuente'], 3)
        self.assertEqual(s['con_gps'], 2)
        self.assertEqual(s['sin_gps'], 1)
        self.assertEqual(s['fuera_radio'], 0)
        self.assertEqual(s['pines'], 2)

    def test_fotos_sin_gps_se_marcan_como_no_gps(self):
        _base, params = self._montar()
        proceso.ejecutar_corrida(params, lambda c: None, lambda m: None)
        wb = load_workbook(params['idx'])
        ws = wb['idxfotos'] if 'idxfotos' in wb.sheetnames else wb[wb.sheetnames[0]]
        # las filas van con el stem en minúsculas (img_7101.jpg) — se compara sin importar mayúsculas
        filas = {str(ws.cell(r, 1).value).lower(): r for r in range(2, ws.max_row + 1)}
        self.assertIn('img_7103.jpg', filas, 'la foto sin GPS no se registró')
        r = filas['img_7103.jpg']
        self.assertEqual(ws.cell(r, 8).value, 'No GPS')          # columna Verificado
        self.assertIn('sin GPS', str(ws.cell(r, 7).value))        # nota en Referencias
        self.assertTrue(ws.cell(r, 9).value in (None, ''))        # sin PIN (no se inventa)
        # las que sí tienen GPS quedan con visto bueno
        r2 = filas['img_7101.jpg']
        self.assertEqual(ws.cell(r2, 8).value, '✓')
        wb.close()

    def test_sin_gps_no_vuelve_a_evaluarse(self):
        _base, params = self._montar()
        proceso.ejecutar_corrida(params, lambda c: None, lambda m: None)
        eventos = []
        res = proceso.ejecutar_corrida(dict(params), eventos.append, lambda m: None)
        s = self._stats(eventos)
        self.assertEqual(s['sin_gps'], 0)      # ya está en el índice -> excluida
        self.assertEqual(res['leidas'], 0)     # las otras dos también

    def test_stats_incluye_capas_del_kml_y_cara(self):
        """PINs por capa (limpiando .xlsx y espacios) + cuántos traen CARA."""
        _base, params = self._montar()
        with open(params['kml'], 'w', encoding='utf-8') as f:
            f.write(KML_DE_PRUEBA)
        eventos = []
        proceso.ejecutar_corrida(params, eventos.append, lambda m: None)
        s = self._stats(eventos)
        self.assertEqual(s['capas'], {'PRUEBA (1)': 2})   # sin '.xlsx'
        self.assertEqual(s['kml'], 2)
        self.assertEqual(s['cara'], 1)                    # solo DF_KML_A tiene CARA


if __name__ == '__main__':
    unittest.main()
