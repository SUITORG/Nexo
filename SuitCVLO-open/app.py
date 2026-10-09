# -*- coding: utf-8 -*-
"""SuitCVLO-open — servidor local (Flask).

Uso: python app.py  ->  http://127.0.0.1:3015
Parámetros: rutas de fotos (local o Drive), salida, pines, KML, idxfotos, N fotos.
Sin base de datos: todo se guarda en archivos locales. Sin borrar archivos.
Lectura IA: ver ia/switch.py (Ollama local -> DeepSeek -> OpenRouter -> Gemini).
"""
import json
import os
import sys
import threading
import uuid
from datetime import datetime

from flask import Flask, jsonify, render_template, request
from werkzeug.exceptions import RequestEntityTooLarge

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import motor
import pdf_reporte
import proceso
import utilidades as U
from ia import switch

BASE = os.path.dirname(os.path.abspath(__file__))
TRABAJO = os.environ.get('CVLO_TRABAJO') or os.path.join(BASE, 'trabajo')
PUERTO = int(os.environ.get('CVLO_PUERTO') or 3015)

app = Flask(__name__,
            template_folder=os.path.join(BASE, 'plantillas'),
            static_folder=os.path.join(BASE, 'static'))
app.config['MAX_CONTENT_LENGTH'] = 1024 * 1024 * 1024

JOBS = {}
JOBS_LOCK = threading.Lock()


def defaults():
    c = r'C:\Users\rojo-\Downloads\cteCVLO1'
    return {
        'salida': c,
        'pines': os.path.join(c, 'GooMapsPines.txt'),
        'drive': os.path.join(c, 'GooDrPics.txt'),
        'kml': os.path.join(c, 'RutaKML.kml'),
        'idx': os.path.join(c, 'idxfotos.xlsx'),
        'drive_url': 'https://drive.google.com/drive/folders/1npZm7uVY0gkqVm8lEcZpaQ7i3P_JVL8l',
    }


def _job_publico(job):
    return {
        'id': job['id'], 'estado': job['estado'], 'modo': job['modo'],
        'total': job['total'], 'completadas': job['completadas'], 'actual': job['actual'],
        'fotos': job['fotos'][-400:], 'log': job['log'][-120:],
        'resumen': job.get('resumen'), 'error': job.get('error_global'),
        'inicio': job['inicio'], 'fin': job.get('fin'),
        'deteniendo': bool(job.get('detener')),
        'stats': job.get('stats'),
    }


@app.get('/')
def inicio():
    return render_template('index.html', defaults_json=json.dumps(defaults()))


@app.get('/api/estado')
def api_estado():
    return jsonify({'mock': motor.en_mock(), 'defaults': defaults(), 'ia': switch.estado()})


@app.get('/api/ia/modelos')
def api_modelos_ia():
    return jsonify({'modelos': switch.modelos(), 'cadena': list(switch.ORDEN)})


@app.get('/api/ia/estado')
def api_estado_ia():
    return jsonify({'proveedores': switch.estado()})


def _ruta_cfg(ruta=None):
    """Archivo de configuración: el elegido en la página > la env CVLO_CFG > CVLOCFG.txt."""
    return (ruta or '').strip() or switch._env('CVLO_CFG') or os.path.join(BASE, 'CVLOCFG.txt')


def _marcas_cfg(cfg):
    """Marcas conocidas del CVLOCFG. Tres formas:

    - `marcas:` **sin valor** + líneas debajo -> lista inline en el propio archivo;
    - `marcas: <ruta>` -> archivo (relativa a esta carpeta);
    - sin nada -> la semilla `marcas-conocidas.txt` del proyecto.
    """
    v = ((cfg or {}).get('marcas') or '').strip()
    if not v:
        v = 'marcas-conocidas.txt'              # semilla (o lista inline vacía: mismas sin marcas)
    if cfg.get('_marcas_inline') or '\n' in v:
        return U.leer_marcas_texto(v)
    if not os.path.isabs(v):
        v = os.path.join(BASE, v)
    return U.leer_marcas(v)


@app.get('/api/config')
def api_config():
    """Rutas precargadas desde el CVLOCFG.txt (opcional).

    `?ruta=` apunta a otro archivo (lo eliges en la página); sin parámetro usa la env
    CVLO_CFG o el CVLOCFG.txt de esta carpeta. Archivo inexistente -> rutas vacías (sin error).
    """
    ruta_cfg = _ruta_cfg(request.args.get('ruta'))
    rutas = U.leer_config_rutas(ruta_cfg)
    return jsonify({'existe': bool(rutas), 'archivo': ruta_cfg, 'rutas': rutas})


def _leer_idx(ruta_idx):
    """(filas, verificadas, no_gps, stems verificados, fotos duplicadas, filas_idx).

    `duplicadas` = fotos que aparecen más de una vez (aunque cambien mayúsculas:
    `IMG_0139.jpg` y `img_0139.jpg` son la misma) — es la "información revuelta".
    `filas_idx` = {stem: {pin, verif, lat, lon}} — para saber de qué campaña es cada
    foto procesada y su distancia al PIN, sin leer el EXIF otra vez.
    """
    vacio = (0, 0, 0, set(), 0, {})
    if not ruta_idx or not os.path.exists(ruta_idx):
        return vacio
    try:
        from openpyxl import load_workbook
        wb = load_workbook(ruta_idx, read_only=True, data_only=True)
        ws = wb['idxfotos'] if 'idxfotos' in wb.sheetnames else wb[wb.sheetnames[0]]
        filas = verif = no_gps = 0
        stems = set()
        vistos = {}
        filas_idx = {}

        def num(v):
            try:
                return float(str(v).replace(',', '.'))
            except Exception:
                return None

        # iter_rows: acceso secuencial (con ws.cell en read_only tardaba 6,5 s)
        for fila in ws.iter_rows(min_row=2, max_col=9, values_only=True):
            nombre = str(fila[0] or '').strip()
            if not nombre:
                continue
            filas += 1
            clave = nombre.lower()
            vistos[clave] = vistos.get(clave, 0) + 1
            v = str(fila[7] or '').strip()
            filas_idx[os.path.splitext(nombre)[0].lower()] = {
                'pin': str(fila[8] or '').strip(),
                'verif': bool(v),
                'lat': num(fila[4]), 'lon': num(fila[5]),
            }
            if v:
                verif += 1
                stems.add(os.path.splitext(nombre)[0].lower())
                if v == 'No GPS':
                    no_gps += 1
        wb.close()
        dup = sum(1 for c in vistos.values() if c > 1)
        return filas, verif, no_gps, stems, dup, filas_idx
    except Exception:
        return vacio


def _campanas(kml_datos, pines_d, filas_idx):
    """Campañas (layers) con sus pines y su avance: pines con >=1 foto leída.

    Las campañas salen del **KML** (los nombres con `.xlsx` del árbol del mapa) — es el
    universo real; si no hay KML, del mapa de pines. **No se unen**: si los dos archivos
    son de mapas distintos, mezclarlos daría pines fantasma.
    `leidos` = pines con >=1 fila en el índice (foto ya leída) → de ahí sale el % de avance.
    `auditados` = pines con >=1 fila con `Verificado ✓` (detalle, no el %).
    """
    fuente = kml_datos or pines_d or {}
    if not fuente:
        return []
    pins_por_capa = {}
    for pin, d in fuente.items():
        pins_por_capa.setdefault(motor.nombre_capa(d), set()).add(pin)
    leidos = set()
    verif_por_pin = {}
    for r in (filas_idx or {}).values():
        if r.get('pin'):
            leidos.add(r['pin'])
            if r.get('verif'):
                verif_por_pin[r['pin']] = verif_por_pin.get(r['pin'], 0) + 1
    res = []
    for nombre, pins in pins_por_capa.items():
        aud = [p for p in pins if p in verif_por_pin]
        res.append({'nombre': nombre, 'pines': len(pins),
                    'leidos': len([p for p in pins if p in leidos]),
                    'auditados': len(aud),
                    'verificadas': sum(verif_por_pin.get(p, 0) for p in pins)})
    res.sort(key=lambda c: (-c['pines'], c['nombre']))
    return res


@app.get('/api/fuentes')
def api_fuentes():
    """Info rápida de las fuentes (sin leer EXIF): fotos, pines, KML e índice."""
    ruta = (request.args.get('ruta') or '').strip()
    ruta_pines = (request.args.get('pines') or '').strip()
    ruta_kml = (request.args.get('kml') or '').strip()
    ruta_idx = (request.args.get('idx') or '').strip()
    salida = (request.args.get('salida') or '').strip()
    out = {'ruta': ruta, 'fotos': None, 'pines': None, 'kml': None, 'indice': None}
    if ruta:
        if not os.path.isdir(ruta):
            return jsonify({'error': 'La ruta de fotos no existe: {}'.format(ruta)}), 400
        out['fotos'] = sum(1 for f in os.listdir(ruta) if U.es_imagen(f))
    pines_d = None
    if ruta_pines:
        try:
            pines_d = motor.cargar_pines(ruta_pines)
            out['pines'] = len(pines_d)
        except Exception as e:
            return jsonify({'error': 'No pude leer los pines: {}'.format(str(e)[:160])}), 400
    kml_datos = None
    if ruta_kml:
        kml_datos = motor.cargar_kml_datos(ruta_kml)
        out['kml'] = motor.kml_resumen(kml_datos)
    filas, verif, no_gps, stems_verif, dup, filas_idx = _leer_idx(ruta_idx)
    out['indice'] = {'filas': filas, 'verificadas': verif, 'no_gps': no_gps, 'duplicadas': dup}
    out['campanas'] = _campanas(kml_datos, pines_d, filas_idx)
    # fotos ya hechas = lo que la corrida ya salta (copiadas o leídas) sobre las de la ruta
    if ruta and os.path.isdir(ruta):
        stems_ruta = {os.path.splitext(f)[0].lower()
                      for f in os.listdir(ruta) if U.es_imagen(f)}
        excl = proceso._excluidos(salida, stems_verif, os.path.join(TRABAJO, 'reads'))
        out['hechas'] = len(excl & stems_ruta)
        out['pendientes'] = len(stems_ruta) - out['hechas']
    if kml_datos and pines_d:
        fuera = len(set(pines_d) - set(kml_datos))
        if fuera:
            # los dos archivos son de mapas distintos: el cruce y las campañas no coincidirán
            out['aviso_mapa'] = ('{} de {} pines del mapa NO están en el KML '
                                 '(¿archivos de mapas distintos?)').format(fuera, len(pines_d))
    return jsonify(out)


@app.get('/api/fuentes/gps')
def api_fuentes_gps():
    """Segundo tiempo: desglose de GPS sobre las fotos NO excluidas (mismos números que la corrida).

    Además calcula, por campaña, la distancia media foto<->pin de las fotos dentro del radio.
    El espejo de la corrida (con/sin/fuera) se calcula contra el mapa de pines; la distancia
    por campaña se atribuye en el **universo de campañas** (el KML si hay) para que coincida
    con la tabla del panel.
    """
    ruta = (request.args.get('ruta') or '').strip()
    ruta_pines = (request.args.get('pines') or '').strip()
    ruta_kml = (request.args.get('kml') or '').strip()
    salida = (request.args.get('salida') or '').strip()
    ruta_idx = (request.args.get('idx') or '').strip()
    try:
        radio = float(request.args.get('radio') or 250)
    except (TypeError, ValueError):
        return jsonify({'error': 'Radio inválido.'}), 400
    if not ruta or not os.path.isdir(ruta):
        return jsonify({'error': 'La ruta de fotos no existe: {}'.format(ruta)}), 400
    if not ruta_pines:
        return jsonify({'error': 'Falta la ruta de GoogleMapsPines.'}), 400
    try:
        pines = motor.cargar_pines(ruta_pines)
    except Exception as e:
        return jsonify({'error': 'No pude leer los pines: {}'.format(str(e)[:160])}), 400

    kml_datos = motor.cargar_kml_datos(ruta_kml) if ruta_kml else {}
    universo = kml_datos or pines          # de dónde salen las campañas de la tabla
    distinto = bool(kml_datos) and set(kml_datos) != set(pines)

    _f, _v, _n, verif, _d, filas_idx = _leer_idx(ruta_idx)
    excl = proceso._excluidos(salida, verif, os.path.join(TRABAJO, 'reads'))
    mock_gps = None
    if motor.en_mock():
        mgp = os.path.join(ruta, 'mock_gps.json')
        if os.path.exists(mgp):
            mock_gps = U.cargar_json(mgp, {})

    # distancia foto<->pin por campaña: procesadas (del índice, sin EXIF) + pendientes
    por_capa = {}

    def acumula(capa, dmete, pendiente):
        e = por_capa.setdefault(capa, {'suma': 0.0, 'n': 0, 'pendientes': 0})
        e['suma'] += dmete
        e['n'] += 1
        if pendiente:
            e['pendientes'] += 1

    for r in filas_idx.values():                      # ya procesadas: foto -> SU pin
        if not r['verif'] or not r['pin'] or r['lat'] is None or r['lon'] is None:
            continue
        d = universo.get(r['pin'])                    # del universo de campañas (KML)
        if not d or 'lat' not in d:
            continue
        acumula(motor.nombre_capa(d), motor.dist_m(r['lat'], r['lon'], d['lat'], d['lon']), False)

    con = sin = fuera = 0
    for f in os.listdir(ruta):
        if not U.es_imagen(f) or proceso._stem(f) in excl:
            continue
        lat, lon, _fecha = motor.gps_y_fecha(os.path.join(ruta, f), mock_gps=mock_gps)
        if lat is None:
            sin += 1
            continue
        con += 1
        if not pines:
            fuera += 1
            continue
        # espejo de la corrida: pin más cercano del mapa de pines
        mejor, dm = None, None
        for pin, d in pines.items():
            if 'lat' not in d:
                continue
            dd = motor.dist_m(lat, lon, d['lat'], d['lon'])
            if dm is None or dd < dm:
                dm, mejor = dd, pin
        if mejor is None or dm > radio:               # fuera del radio: no cruza
            fuera += 1
        # distancia por campaña: en el universo de campañas (si es distinto del mapa)
        if distinto:
            pin_u, dm_u = None, None
            for pin, d in universo.items():
                if 'lat' not in d:
                    continue
                dd = motor.dist_m(lat, lon, d['lat'], d['lon'])
                if dm_u is None or dd < dm_u:
                    dm_u, pin_u = dd, pin
            if pin_u is not None and dm_u <= radio:
                acumula(motor.nombre_capa(universo[pin_u]), dm_u, True)
        elif mejor is not None and dm <= radio:
            acumula(motor.nombre_capa(pines[mejor]), dm, True)

    out = {'con_gps': con, 'sin_gps': sin, 'fuera_radio': fuera, 'radio': radio}
    out['campanas'] = {c: {'dist_prom': (round(v['suma'] / v['n'], 1) if v['n'] else None),
                           'n': v['n'], 'pendientes': v['pendientes']}
                       for c, v in por_capa.items()}
    return jsonify(out)


def _params_repro():
    """query string -> params para proceso. Devuelve (params, None) o (None, (respuesta, 400))."""
    ruta = (request.args.get('ruta') or '').strip()
    salida = (request.args.get('salida') or '').strip()
    if not ruta or not os.path.isdir(ruta):
        return None, (jsonify({'error': 'La ruta de fotos no existe: {}'.format(ruta)}), 400)
    if not salida:
        return None, (jsonify({'error': 'Falta la Ruta de Salida.'}), 400)
    return {'ruta_local': ruta, 'salida': salida,
            'idx': (request.args.get('idx') or '').strip(),
            'trabajo': TRABAJO}, None


@app.get('/api/reprocesar/preview')
def api_repro_preview():
    """Conteo de fotos incompletas — sin leer EXIF, <1 s."""
    params, err = _params_repro()
    if err:
        return err
    obj, sin_origen, sin_lectura, viejo = proceso._objetivos_reproceso(params)
    return jsonify({'objetivas': len(obj), 'sin_lectura': len(sin_lectura),
                    'formato_viejo': len(viejo), 'sin_origen': len(sin_origen)})


@app.post('/api/reprocesar')
def api_repro():
    """Respalda el idxfotos.xlsx y limpia 'Verificado' de las fotos incompletas."""
    params, err = _params_repro()
    if err:
        return err
    try:
        res = proceso.preparar_reproceso(params, lambda m: app.logger.info('[reproceso] %s', m))
    except Exception as e:
        return jsonify({'error': str(e)[:300]}), 400
    if not res['objetivas']:
        return jsonify({'error': 'Nada que reprocesar: todo ya está al día.'}), 400
    return jsonify(res)


@app.errorhandler(RequestEntityTooLarge)
def _grande(_e):
    return jsonify({'error': 'La carga excede el tamaño máximo (1 GB).'}), 413


@app.errorhandler(Exception)
def _error_global(e):
    """Cualquier error inesperado responde JSON en vez de un 500 sin detalle."""
    from werkzeug.exceptions import HTTPException
    if isinstance(e, HTTPException):
        return e
    app.logger.exception('Error en %s %s', request.method, request.path)
    return jsonify({'error': 'Error interno: {}'.format(str(e)[:300])}), 500


@app.post('/api/corridas')
def api_crear():
    form = request.form
    modo = (form.get('modo') or 'local').strip().lower()
    try:
        n_fotos = int(form.get('n_fotos') or 50)
        radio = float(form.get('radio') or 250)
        cap = int(form.get('cap') or 10)
        pausa = float(form.get('pausa') or 1.0)
        max_descargas = int(form.get('max_descargas') or 300)
    except (TypeError, ValueError):
        return jsonify({'error': 'Parámetros numéricos inválidos (N fotos, radio, tope o pausa).'}), 400
    salida = (form.get('salida') or '').strip()
    pines = (form.get('pines') or '').strip()
    kml = (form.get('kml') or '').strip()
    idx = (form.get('idx') or '').strip()
    drive = (form.get('drive_pics') or '').strip()
    ruta_local = (form.get('ruta_local') or '').strip()
    ia_proveedor = (form.get('ia_proveedor') or 'auto').strip().lower()
    ia_modelo = (form.get('ia_modelo') or '').strip()
    comentario = (form.get('comentario') or '1').strip().lower() not in ('0', 'false', 'no')
    # prioridad de la fuente: la página (select) manda; si no viene, la del CVLOCFG; si no, kml
    cfg = U.leer_config_rutas(_ruta_cfg(form.get('cfg_ruta')))
    prioridad = U.prioridad_valida(form.get('prioridad') or cfg.get('prioridad'))

    if not salida:
        return jsonify({'error': 'Falta la Ruta de Salida.'}), 400
    if not pines:
        return jsonify({'error': 'Falta la Ruta de GoogleMapsPines.'}), 400
    if not idx:
        return jsonify({'error': 'Falta la Ruta de idxfotos.xlsx.'}), 400
    if ia_proveedor not in ('auto',) + tuple(switch.ORDEN):
        return jsonify({'error': 'Motor IA desconocido: {}'.format(ia_proveedor)}), 400

    jid = uuid.uuid4().hex[:12]
    switch.reset()   # nueva corrida: no arrastrar los cooldowns de la anterior
    entrada = os.path.join(TRABAJO, 'corridas', jid)
    os.makedirs(entrada, exist_ok=True)

    # --- fuente de fotos ---
    if modo in ('zip', 'carpeta'):
        destino = os.path.join(entrada, 'fotos')
        os.makedirs(destino, exist_ok=True)
        archivo_zip = request.files.get('archivo_zip')
        if archivo_zip is not None and archivo_zip.filename:
            ruta_zip = os.path.join(entrada, U.sanear_nombre(archivo_zip.filename, 'lote.zip'))
            archivo_zip.save(ruta_zip)
            try:
                U.extraer_zip_seguro(ruta_zip, destino)
            except Exception as e:
                return jsonify({'error': 'ZIP inválido: {}'.format(str(e)[:150])}), 400
        for f in request.files.getlist('archivos'):
            if f is None or not f.filename:
                continue
            nombre = U.nombre_limpio(f.filename)
            if not U.es_imagen(nombre):
                continue
            f.save(U.ruta_libre(destino, nombre))
        ruta_local = destino
    elif modo == 'web':
        # aceptar URL directa o un .txt que la contenga (GooDrPics.txt)
        u = drive
        if u and os.path.isdir(u):
            return jsonify({'error': 'Modo Web: "{}" es una CARPETA. Indica la ruta del archivo '
                                     'GooDrPics.txt (o pega la URL de Drive).'.format(u)}), 400
        if u and os.path.isfile(u):
            try:
                texto = U.leer_texto_auto(u)
            except Exception as e:
                return jsonify({'error': 'No pude leer "{}": {}'.format(u, str(e)[:150])}), 400
            import re
            m = re.search(r'https?://\S+', texto)
            u = m.group(0) if m else ''
        if not u or 'http' not in u:
            return jsonify({'error': 'Modo Web: indica una URL de Drive válida (o la ruta del archivo GooDrPics.txt).'}), 400
        drive = u
        if motor.en_mock():
            return jsonify({'error': 'MODO_MOCK: la fuente web no está disponible en pruebas.'}), 400
    else:
        if not ruta_local:
            return jsonify({'error': 'Modo Local: falta la ruta de fotos.'}), 400
        if os.path.isfile(ruta_local):
            return jsonify({'error': 'Modo Local: "{}" es un archivo. Indica la CARPETA que contiene las fotos.'
                           .format(ruta_local)}), 400
        if not os.path.isdir(ruta_local):
            return jsonify({'error': 'Modo Local: la ruta de fotos no existe: {}'.format(ruta_local)}), 400
        modo = 'local'

    mock_gps = None
    if motor.en_mock() and ruta_local:
        mgp = os.path.join(ruta_local, 'mock_gps.json')
        if os.path.exists(mgp):
            mock_gps = U.cargar_json(mgp, {})

    job = {
        'id': jid, 'modo': modo, 'estado': 'en cola', 'total': 0, 'completadas': 0, 'actual': '',
        'fotos': [], 'log': [], 'resumen': None, 'error_global': None, 'detener': False,
        'inicio': datetime.now().isoformat(timespec='seconds'), 'fin': None,
    }
    with JOBS_LOCK:
        JOBS[jid] = job

    params = {
        'modo': modo, 'ruta_local': ruta_local, 'salida': salida, 'pines': pines,
        'drive_pics': drive, 'kml': kml, 'idx': idx, 'n_fotos': n_fotos,
        'radio': radio, 'cap': cap, 'pausa': pausa, 'max_descargas': max_descargas,
        'comentario': comentario, '_reproceso': (form.get('reproceso') == '1'),
        'prioridad': prioridad, '_marcas': _marcas_cfg(cfg),
        'ia_proveedor': ia_proveedor, 'ia_modelo': ia_modelo,
        'trabajo': TRABAJO,
        '_mock_gps': mock_gps,
        '_detener': lambda: bool(job.get('detener')),
    }

    def progreso(cambios):
        with JOBS_LOCK:
            job.update(cambios)

    def log(mensaje):
        with JOBS_LOCK:
            job['log'].append('{}  {}'.format(datetime.now().strftime('%H:%M:%S'), mensaje))
            job['log'] = job['log'][-300:]

    def correr():
        try:
            resumen = proceso.ejecutar_corrida(params, progreso, log)
            resumen['grupos'] = list(resumen.get('grupos') or [])
            resumen['pines'] = list(resumen.get('pines') or [])
            progreso({'estado': 'listo', 'resumen': resumen, 'fin': datetime.now().isoformat(timespec='seconds')})
            log('Corrida terminada.')
        except Exception as e:
            progreso({'estado': 'error', 'error_global': str(e)[:400],
                      'fin': datetime.now().isoformat(timespec='seconds')})
            log('ERROR: {}'.format(str(e)[:250]))

    threading.Thread(target=correr, name='cvlo-' + jid, daemon=True).start()
    return jsonify({'id': jid})


@app.get('/api/corridas/<jid>')
def api_job(jid):
    with JOBS_LOCK:
        job = JOBS.get(jid)
    if not job:
        return jsonify({'error': 'no existe'}), 404
    return jsonify(_job_publico(job))


@app.post('/api/corridas/<jid>/detener')
def api_detener(jid):
    with JOBS_LOCK:
        job = JOBS.get(jid)
        if not job:
            return jsonify({'error': 'no existe'}), 404
        if job['estado'] in ('listo', 'error'):
            return jsonify({'estado': job['estado'], 'aviso': 'La corrida ya terminó.'})
        job['detener'] = True
        job['log'].append('{}  Detener solicitado: se corta al terminar la foto en curso.'
                          .format(datetime.now().strftime('%H:%M:%S')))
    return jsonify({'estado': 'deteniendo'})


@app.get('/api/corridas')
def api_lista():
    with JOBS_LOCK:
        items = [{'id': j['id'], 'modo': j['modo'], 'estado': j['estado'], 'inicio': j['inicio'],
                  'completadas': j['completadas'], 'total': j['total']} for j in JOBS.values()]
    items.sort(key=lambda x: x['inicio'], reverse=True)
    return jsonify({'corridas': items[:20]})


@app.post('/api/marcas/actualizar')
def api_marcas_actualizar():
    """Alimenta la lista de marcas: agrega a la fuente en uso las vistas en las lecturas.

    Fuente en uso = lista inline del CFG > `marcas: <ruta>` > `marcas-conocidas.txt`.
    No llama a la IA: lee `trabajo/reads/*.json` (mismo criterio que la columna `Marca`).
    """
    form = request.form
    ruta_cfg = _ruta_cfg(form.get('cfg_ruta'))
    cfg = U.leer_config_rutas(ruta_cfg)
    vistas = proceso.marcas_vistas(os.path.join(TRABAJO, 'reads'))
    tipo, ruta = U.marcas_destino(ruta_cfg, cfg)
    res = U.marcas_agregar(tipo, ruta, vistas)
    res['vistas'] = len(vistas)
    res['tipo'] = tipo
    return jsonify(res)


@app.post('/api/fuentes/pdf')
def api_fuentes_pdf():
    """Escribe el análisis de fuentes como PDF en la Ruta de Salida (fecha y hora en el nombre).

    Recibe del navegador las líneas ya calculadas (mismas que ve en pantalla) y las maquilla
    con la cabecera: título + `Generado: AAAA-MM-DD HH:MM`. Sin dependencias: pdf_reporte.
    """
    data = request.get_json(silent=True) or {}
    salida = (data.get('salida') or '').strip()
    if not salida:
        return jsonify({'error': 'Falta la Ruta de Salida (ahí se guarda el PDF).'}), 400
    if not os.path.isdir(salida):
        return jsonify({'error': 'La Ruta de Salida no existe: {}'.format(salida)}), 400
    lineas = data.get('lineas')
    if not isinstance(lineas, list) or not lineas:
        return jsonify({'error': 'Sin contenido de análisis.'}), 400
    ahora = datetime.now()
    cabecera = ['ANALISIS DE FUENTES - SuitCVLO-open',
                'Generado: ' + ahora.strftime('%Y-%m-%d %H:%M')]
    nombre = 'analisis_fuentes_' + ahora.strftime('%Y-%m-%d_%H-%M') + '.pdf'
    destino = os.path.join(salida, nombre)
    try:
        ruta, nbytes = pdf_reporte.generar_pdf(cabecera + [str(l) for l in lineas], destino)
    except Exception as e:
        return jsonify({'error': 'No pude generar el PDF: {}'.format(str(e)[:250])}), 500
    return jsonify({'archivo': nombre, 'ruta': ruta, 'bytes': nbytes})


@app.post('/api/reescribir-info')
def api_reescribir_info():
    """Reescribe TODOS los info.txt de la salida con las lecturas ya guardadas (sin IA).

    Sirve para aplicar un cambio de prioridad (u otro ajuste del bloque) al entregable
    existente: no lee ninguna foto, no copia ni borra archivos.
    """
    form = request.form
    salida = (form.get('salida') or '').strip()
    if not salida:
        return jsonify({'error': 'Falta la Ruta de Salida.'}), 400
    if not os.path.isdir(salida):
        return jsonify({'error': 'La Ruta de Salida no existe: {}'.format(salida)}), 400
    cfg = U.leer_config_rutas(_ruta_cfg(form.get('cfg_ruta')))
    prioridad = U.prioridad_valida(form.get('prioridad') or cfg.get('prioridad'))
    comentario = (form.get('comentario') or '1').strip().lower() not in ('0', 'false', 'no')
    ruta_pines = (form.get('pines') or '').strip()
    if ruta_pines:
        try:
            pdatos = motor.cargar_pines(ruta_pines)
        except Exception as e:
            return jsonify({'error': 'No pude leer los pines: {}'.format(str(e)[:150])}), 400
    else:
        pdatos = {}
    kdatos = motor.cargar_kml_datos((form.get('kml') or '').strip())
    marcas = _marcas_cfg(cfg)
    res = proceso.reescribir_todo(salida, pdatos, kdatos, os.path.join(TRABAJO, 'reads'),
                                  prioridad=prioridad, marcas=marcas, comentario=comentario,
                                  idx=(form.get('idx') or '').strip())
    return jsonify(res)


if __name__ == '__main__':
    print('SuitCVLO-open')
    print('  URL: http://127.0.0.1:{}'.format(PUERTO))
    print('  MODO_MOCK: {}'.format('ACTIVADO' if motor.en_mock() else 'desactivado'))
    if not os.environ.get('CVLO_NO_ABRIR'):
        import webbrowser
        threading.Timer(1.5, lambda: webbrowser.open('http://127.0.0.1:{}'.format(PUERTO))).start()
    app.run(host='127.0.0.1', port=PUERTO, threaded=True, use_reloader=False)
