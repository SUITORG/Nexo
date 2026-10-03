# -*- coding: utf-8 -*-
"""E2E en modo MOCK: proceso completo + API de la app."""
import json
import os
import sys
import tempfile
import time
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

os.environ['CVLO_MOCK'] = '1'

from PIL import Image
from openpyxl import load_workbook

import motor
import proceso


def crear_fuente(base, fotos_datos):
    """fotos_datos: [(nombre, fecha, lat, lon)] + mock_gps.json + pines.xlsx."""
    fotos = os.path.join(base, 'fotos')
    os.makedirs(fotos, exist_ok=True)
    mock = {}
    for i, (nombre, fecha, lat, lon) in enumerate(fotos_datos):
        Image.new('RGB', (320, 240), (240 - i * 10, 240, 240)).save(os.path.join(fotos, nombre), quality=90)
        mock[nombre] = {'lat': lat, 'lon': lon, 'fecha': fecha}
    with open(os.path.join(fotos, 'mock_gps.json'), 'w', encoding='utf-8') as f:
        json.dump(mock, f)
    # pines
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Pines'
    ws.append(['PIN', 'LATITUD', 'LONGITUD', 'MEDIO', 'FORMATO', 'ID', 'UBICACIÓN', 'CIUDAD', 'CAPA'])
    ws.append(['DF_09_TEST_1', 25.7000, -100.3500, 'PARABUS', 'FIJO', 'T-1', 'CALLE UNO', 'MTY', 'DIDI PRUEBA'])
    ws.append(['DF_09_TEST_2', 25.7100, -100.3600, 'BILLBOARD', 'FIJO', 'T-2', 'CALLE DOS', 'MTY', 'DIDI PRUEBA'])
    pines = os.path.join(base, 'pines.xlsx')
    wb.save(pines)
    kml = os.path.join(base, 'ruta.kml')
    open(kml, 'w', encoding='utf-8').write('<kml/>')
    return fotos, pines, kml


class TestFlujoMock(unittest.TestCase):

    def test_corrida_completa(self):
        base = tempfile.mkdtemp(prefix='cvlo_e2e_')
        fotos, pines, kml = crear_fuente(base, [
            ('IMG_0001.jpg', '2026-09-14', 25.70002, -100.3500),
            ('IMG_0002.jpg', '2026-09-14', 25.70006, -100.3500),
            ('IMG_0003.jpg', '2026-09-20', 25.71002, -100.3600)])
        salida = os.path.join(base, 'salida')
        idx = os.path.join(base, 'idxfotos.xlsx')
        trabajo = os.path.join(base, 'trabajo')

        params = {'modo': 'local', 'ruta_local': fotos, 'salida': salida, 'pines': pines,
                  'drive_pics': '', 'kml': kml, 'idx': idx, 'n_fotos': 10, 'radio': 250,
                  'cap': 10, 'pausa': 0, 'trabajo': trabajo,
                  '_mock_gps': json.load(open(os.path.join(fotos, 'mock_gps.json'), encoding='utf-8'))}
        eventos = []
        resumen = proceso.ejecutar_corrida(params, lambda c: eventos.append(c), lambda m: eventos.append(m))

        self.assertEqual(resumen['leidas'], 3)
        self.assertEqual(resumen['copiadas'], 3)
        # estructura 260901 y 260902 (según fecha de foto)
        f1 = os.path.join(salida, '260901', 'auditoria', 'DIDI PARABUSES MTY', 'DF_09_TEST_1')
        f3 = os.path.join(salida, '260902', 'auditoria', 'DIDI PANORAMICOS MTY', 'DF_09_TEST_2')
        self.assertTrue(os.path.exists(os.path.join(f1, 'IMG_0001.jpg')))
        self.assertTrue(os.path.exists(os.path.join(f1, 'IMG_0002.jpg')))
        self.assertTrue(os.path.exists(os.path.join(f3, 'IMG_0003.jpg')))
        self.assertTrue(os.path.exists(os.path.join(f1, 'info.txt')))
        info = open(os.path.join(f1, 'info.txt'), encoding='utf-8').read()
        self.assertIn('PIN: DF_09_TEST_1', info)
        self.assertIn('CAMPAÑA:', info)
        # idxfotos creado con 3 filas verificadas
        wb = load_workbook(idx)
        ws = wb['idxfotos']
        self.assertEqual(ws.max_row - 1, 3)
        verif = [ws.cell(r, 8).value for r in range(2, ws.max_row + 1)]
        self.assertEqual(sum(1 for v in verif if v), 3)

        # segunda corrida: no repite nada
        params2 = dict(params)
        params2['_mock_gps'] = params['_mock_gps']
        resumen2 = proceso.ejecutar_corrida(params2, lambda c: None, lambda m: None)
        self.assertEqual(resumen2['leidas'], 0)

    def test_tope_por_pin(self):
        base = tempfile.mkdtemp(prefix='cvlo_cap_')
        fotos, pines, kml = crear_fuente(base, [
            ('IMG_1001.jpg', '2026-09-14', 25.70002, -100.3500),
            ('IMG_1002.jpg', '2026-09-14', 25.70004, -100.3500),
            ('IMG_1003.jpg', '2026-09-14', 25.70006, -100.3500)])
        salida = os.path.join(base, 'salida')
        idx = os.path.join(base, 'idxfotos.xlsx')
        params = {'modo': 'local', 'ruta_local': fotos, 'salida': salida, 'pines': pines,
                  'drive_pics': '', 'kml': kml, 'idx': idx, 'n_fotos': 10, 'radio': 250,
                  'cap': 2, 'pausa': 0, 'trabajo': os.path.join(base, 'trabajo'),
                  '_mock_gps': json.load(open(os.path.join(fotos, 'mock_gps.json'), encoding='utf-8'))}
        resumen = proceso.ejecutar_corrida(params, lambda c: None, lambda m: None)
        self.assertEqual(resumen['leidas'], 3)
        self.assertEqual(resumen['copiadas'], 2)
        self.assertEqual(resumen['excedidas'], 1)
        wb = load_workbook(idx)
        ws = wb['idxfotos']
        blancos = [r for r in range(2, ws.max_row + 1) if not ws.cell(r, 8).value]
        self.assertEqual(len(blancos), 1)
        nota = str(ws.cell(blancos[0], 7).value or '')
        self.assertIn('excede', nota)


class TestFechaCarga(unittest.TestCase):
    """Columna extra que pidió el usuario: fecha en que se cargó cada foto."""

    def test_xlsx_nuevo_lleva_fecha_carga(self):
        base = tempfile.mkdtemp(prefix='cvlo_fecha_')
        fotos, pines, kml = crear_fuente(base, [('IMG_3001.jpg', '2026-09-14', 25.70002, -100.3500)])
        params = {'modo': 'local', 'ruta_local': fotos, 'salida': os.path.join(base, 'salida'),
                  'pines': pines, 'drive_pics': '', 'kml': kml,
                  'idx': os.path.join(base, 'idxfotos.xlsx'), 'n_fotos': 10, 'radio': 250,
                  'cap': 10, 'pausa': 0, 'trabajo': os.path.join(base, 'trabajo'),
                  '_mock_gps': json.load(open(os.path.join(fotos, 'mock_gps.json'), encoding='utf-8'))}
        proceso.ejecutar_corrida(params, lambda c: None, lambda m: None)
        wb = load_workbook(params['idx'])
        ws = wb['idxfotos']
        self.assertEqual(ws.cell(1, 9).value, 'PIN')
        self.assertEqual(ws.cell(1, 10).value, 'FechaCarga')
        self.assertRegex(str(ws.cell(2, 10).value or ''), r'^\d{4}-\d{2}-\d{2} \d{2}:\d{2}$')

    def test_xlsx_viejo_gana_la_columna_sin_tocar_las_demas(self):
        from openpyxl import Workbook
        base = tempfile.mkdtemp(prefix='cvlo_viejo_')
        ruta = os.path.join(base, 'idxfotos.xlsx')
        wb = Workbook()
        ws = wb.active
        ws.title = 'idxfotos'
        ws.append(['NombreFoto', 'Descripción de lo que anuncia', 'Dirección', 'Tipo de Anuncio',
                   'Latitud', 'Longitud', 'Referencias visuales o notas', 'Verificado', 'PIN'])
        ws.append(['vieja.jpg', 'que anuncia', 'calle', 'mupi', 1, 2, 'refs', 'ok', 'DF_1'])
        wb.save(ruta)

        wb2, ws2 = motor.idx_abrir_o_crear(ruta)
        self.assertEqual(ws2.cell(1, 10).value, 'FechaCarga')
        self.assertEqual(ws2.cell(1, 9).value, 'PIN')
        self.assertEqual(ws2.cell(2, 1).value, 'vieja.jpg')
        wb2.save(ruta)

    def test_no_pisa_una_columna_de_fecha_ajena(self):
        """El usuario ya tiene 'FechaHoraFoto' en la columna 10: FechaCarga va a la 11."""
        from openpyxl import Workbook
        base = tempfile.mkdtemp(prefix='cvlo_ajena_')
        ruta = os.path.join(base, 'idxfotos.xlsx')
        wb = Workbook()
        ws = wb.active
        ws.title = 'Sheet1'                      # ni siquiera se llama idxfotos
        ws.append(['NombreFoto', 'Descripción de lo que anuncia', 'Dirección', 'Tipo de Anuncio',
                   'Latitud', 'Longitud', 'Referencias visuales o notas', 'Verificado', 'PIN',
                   'FechaHoraFoto'])
        ws.append(['mia.jpg', 'que', 'calle', 'mupi', 1, 2, 'refs', 'ok', 'DF_2', '2026-10-01 08:00'])
        wb.save(ruta)

        wb2, ws2 = motor.idx_abrir_o_crear(ruta)
        col = motor.idx_col_fecha(ws2)
        self.assertEqual(col, 11)
        self.assertEqual(ws2.cell(1, 10).value, 'FechaHoraFoto')      # intacta
        self.assertEqual(ws2.cell(2, 10).value, '2026-10-01 08:00')   # sus datos, intactos
        self.assertEqual(ws2.cell(1, 11).value, 'FechaCarga')
        wb2.save(ruta)

        # y la corrida escribe la fecha en la 11, sin tocar la 10
        fotos, pines, kml = crear_fuente(base, [('IMG_4001.jpg', '2026-09-14', 25.70002, -100.3500)])
        params = {'modo': 'local', 'ruta_local': fotos, 'salida': os.path.join(base, 'salida'),
                  'pines': pines, 'drive_pics': '', 'kml': kml, 'idx': ruta,
                  'n_fotos': 10, 'radio': 250, 'cap': 10, 'pausa': 0,
                  'trabajo': os.path.join(base, 'trabajo'),
                  '_mock_gps': json.load(open(os.path.join(fotos, 'mock_gps.json'), encoding='utf-8'))}
        proceso.ejecutar_corrida(params, lambda c: None, lambda m: None)
        wb3 = load_workbook(ruta)
        ws3 = wb3[wb3.sheetnames[0]]
        self.assertEqual(ws3.cell(2, 10).value, '2026-10-01 08:00')
        self.assertRegex(str(ws3.cell(3, 11).value or ''), r'^\d{4}-\d{2}-\d{2} \d{2}:\d{2}$')


    def test_info_txt_lleva_comentario(self):
        """Cada bloque del info.txt lleva COMENTARIO con el mensaje del anuncio."""
        base = tempfile.mkdtemp(prefix='cvlo_coment_')
        fotos, pines, kml = crear_fuente(base, [('IMG_6201.jpg', '2026-09-14', 25.70002, -100.3500)])
        params = {'modo': 'local', 'ruta_local': fotos, 'salida': os.path.join(base, 'salida'),
                  'pines': pines, 'drive_pics': '', 'kml': kml,
                  'idx': os.path.join(base, 'idxfotos.xlsx'), 'n_fotos': 10, 'radio': 250,
                  'cap': 10, 'pausa': 0, 'trabajo': os.path.join(base, 'trabajo'),
                  '_mock_gps': json.load(open(os.path.join(fotos, 'mock_gps.json'), encoding='utf-8'))}
        proceso.ejecutar_corrida(params, lambda c: None, lambda m: None)
        ruta = os.path.join(base, 'salida', '260901', 'auditoria',
                            'DIDI PARABUSES MTY', 'DF_09_TEST_1', 'info.txt')
        self.assertTrue(os.path.exists(ruta), ruta)
        texto = open(ruta, encoding='utf-8').read()
        self.assertIn('COMENTARIO:', texto)
        self.assertIn('promoción simulada', texto)
        self.assertLess(texto.index('COMENTARIO:'), texto.index('COMENTARIOS:'))

    def test_rellena_fecha_hora_subida_sin_crear_columna_nueva(self):
        """Su xlsx real: col10 FechaHoraFoto, col11 FechaHoraSubida → se rellena la 11."""
        from openpyxl import Workbook
        base = tempfile.mkdtemp(prefix='cvlo_subida_')
        ruta = os.path.join(base, 'idxfotos.xlsx')
        wb = Workbook()
        ws = wb.active
        ws.title = 'Sheet1'
        ws.append(['NombreFoto', 'Descripción de lo que anuncia', 'Dirección', 'Tipo de Anuncio',
                   'Latitud', 'Longitud', 'Referencias visuales o notas', 'Verificado', 'PIN',
                   'FechaHoraFoto', 'FechaHoraSubida'])
        ws.append(['vieja.jpg', 'que anuncia', 'calle', 'mupi', 1, 2, 'refs', 'ok',
                   'DF_09_TEST_1', '2026-09-01 08:00', '2026-09-01 09:00'])
        wb.save(ruta)

        fotos, pines, kml = crear_fuente(base, [('IMG_6301.jpg', '2026-09-14', 25.70002, -100.3500)])
        params = {'modo': 'local', 'ruta_local': fotos, 'salida': os.path.join(base, 'salida'),
                  'pines': pines, 'drive_pics': '', 'kml': kml, 'idx': ruta,
                  'n_fotos': 10, 'radio': 250, 'cap': 10, 'pausa': 0,
                  'trabajo': os.path.join(base, 'trabajo'),
                  '_mock_gps': json.load(open(os.path.join(fotos, 'mock_gps.json'), encoding='utf-8'))}
        proceso.ejecutar_corrida(params, lambda c: None, lambda m: None)

        wb2 = load_workbook(ruta)
        ws2 = wb2[wb2.sheetnames[0]]
        self.assertEqual(ws2.max_column, 11)                     # no inventó columnas
        self.assertEqual(ws2.cell(1, 11).value, 'FechaHoraSubida')
        self.assertEqual(ws2.cell(2, 11).value, '2026-09-01 09:00')
        nueva = ws2.max_row
        self.assertRegex(str(ws2.cell(nueva, 11).value or ''),
                         r'^\d{4}-\d{2}-\d{2} \d{2}:\d{2}$')
        wb2.close()


if __name__ == '__main__':
    unittest.main()
